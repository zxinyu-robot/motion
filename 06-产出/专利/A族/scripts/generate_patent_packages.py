#!/usr/bin/env python3
"""Batch-generate PNG figures and Word disclosure docs for A-family patents."""

from __future__ import annotations

import re
import subprocess
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent.parent
FIG_DIR = BASE / "附图"
SCRIPTS = Path(__file__).resolve().parent


@dataclass
class PatentCase:
    key: str
    md_name: str
    fig_script: str
    figures: list[str] = field(default_factory=list)

    @property
    def md_path(self) -> Path:
        return BASE / self.md_name

    @property
    def docx_path(self) -> Path:
        return self.md_path.with_suffix(".docx")


CASES: list[PatentCase] = [
    PatentCase(
        "A1",
        "01_发明专利_多机异构空间数据运行时架构_技术交底书.md",
        "fig_p01.py",
        [
            "01-图1-系统总体架构图.png",
            "01-图2-机器人端数据采集流程图.png",
            "01-图3-边缘端子图融合流程图.png",
            "01-图4-统一数据对象关系图.png",
            "01-图5-世界模型更新与查询流程图.png",
            "01-图6-多场景实施例示意图.png",
            "01-图7-自组网拓扑与语义优先分级回传图.png",
        ],
    ),
    PatentCase(
        "A2",
        "02_发明专利_基于地图质量的子图自适应上传融合_技术交底书.md",
        "fig_p02.py",
        [
            "02-图1-子图自适应上传方法总体流程图.png",
            "02-图2-地图质量评分计算模块图.png",
            "02-图3-多级上传策略状态转换图.png",
            "02-图4-边缘端融合队列处理流程图.png",
            "02-图5-弱网环境摘要上传与补传流程图.png",
        ],
    ),
    PatentCase(
        "A3",
        "03_发明专利_WorldModel增量更新与查询方法_技术交底书.md",
        "fig_p03.py",
        [
            "03-图1-世界模型总体层级结构图.png",
            "03-图2-世界模型增量更新流程图.png",
            "03-图3-多源观测冲突处理流程图.png",
            "03-图4-查询接口与上层应用关系图.png",
            "03-图5-世界模型与各系统交互图.png",
        ],
    ),
    PatentCase(
        "A4",
        "04_发明专利_语义优先与自组网自适应的空间数据语义化上传融合_技术交底书.md",
        "fig_p04.py",
        [
            "04-图1-方法总体流程图.png",
            "04-图2-移动自组网拓扑与节点角色切换示意图.png",
            "04-图3-语义价值分级与多级语义包结构图.png",
            "04-图4-自组网链路自适应上传决策模块图.png",
            "04-图5-边缘端语义解码与可信融合流程图.png",
            "04-图6-机会式缓存与补传流程图.png",
            "04-图7-TCP与数据链路层语义传输协议栈对比图.png",
            "04-图8-多模异构链路自适应承载与协议映射图.png",
        ],
    ),
]


def generate_figures(case: PatentCase) -> None:
    script = FIG_DIR / case.fig_script
    if not script.exists():
        raise FileNotFoundError(f"Missing figure script: {script}")
    subprocess.run([sys.executable, str(script)], check=True, cwd=FIG_DIR)


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text.strip(), level=min(level, 3))
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = row[j] if j < len(row) else ""
            cell_text = re.sub(r"\*\*(.+?)\*\*", r"\1", cell_text)
            cell_text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", cell_text)
            cell_text = re.sub(r"`(.+?)`", r"\1", cell_text)
            table.rows[i].cells[j].text = cell_text


def add_figure(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.5))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.font.size = Pt(10.5)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_paragraph(doc: Document, text: str) -> None:
    text = text.strip()
    if not text:
        return
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_bullet(doc: Document, text: str) -> None:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text.strip())
    text = re.sub(r"`(.+?)`", r"\1", text)
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_numbered(doc: Document, text: str) -> None:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text.strip())
    text = re.sub(r"`(.+?)`", r"\1", text)
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def build_docx(md_path: Path, md_text: str) -> Document:
    doc = Document()
    set_doc_font(doc)

    lines = md_text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith(">"):
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("##"):
            title = stripped.lstrip("# ").strip()
            t = doc.add_paragraph()
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = t.add_run(title)
            run.bold = True
            run.font.size = Pt(16)
            run.font.name = "黑体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            i += 1
            continue

        if stripped.startswith("## "):
            add_heading(doc, stripped[3:], 1)
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 2)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table_rows(table_lines))
            continue

        if re.match(r"^\d+\.\s", stripped):
            add_numbered(doc, re.sub(r"^\d+\.\s*", "", stripped))
            i += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            i += 1
            continue

        if stripped.startswith("![") and "](" in stripped:
            m = re.match(r"!\[(.+?)\]\((.+?)\)", stripped)
            if m:
                caption, rel = m.group(1), m.group(2)
                img_path = (md_path.parent / rel).resolve()
                if img_path.exists():
                    add_figure(doc, img_path, caption)
            i += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and len(stripped) < 120:
            p = doc.add_paragraph(stripped.strip("*"))
            for run in p.runs:
                run.bold = True
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            continue

        if stripped:
            add_paragraph(doc, stripped)

        i += 1

    return doc


def generate_docx(case: PatentCase) -> Path:
    md_path = case.md_path
    if not md_path.exists():
        raise FileNotFoundError(f"Missing markdown: {md_path}")

    for name in case.figures:
        png = FIG_DIR / name
        if not png.exists():
            raise FileNotFoundError(f"Missing PNG figure: {png}")

    md_text = md_path.read_text(encoding="utf-8")
    doc = build_docx(md_path, md_text)
    doc.save(case.docx_path)
    print(f"DOCX: {case.docx_path.relative_to(BASE)}")
    return case.docx_path


def create_zip(cases: list[PatentCase], extra_files: list[str]) -> Path:
    stamp = date.today().strftime("%Y%m%d")
    zip_path = BASE / f"A族-A1A4-{stamp}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for case in cases:
            zf.write(case.docx_path, case.docx_path.name)
            for fig in case.figures:
                png = FIG_DIR / fig
                zf.write(png, f"附图/{png.name}")
        for name in extra_files:
            path = BASE / name
            if path.exists():
                zf.write(path, path.name)
    print(f"ZIP: {zip_path.relative_to(BASE)}")
    return zip_path


def main(argv: list[str]) -> int:
    only_keys = [a for a in argv if not a.startswith("-")]
    png_only = "--png-only" in argv
    docx_only = "--docx-only" in argv
    no_zip = "--no-zip" in argv

    selected = CASES
    if only_keys:
        selected = [c for c in CASES if c.key in only_keys]
        if not selected:
            print(f"Unknown keys: {only_keys}. Valid: {[c.key for c in CASES]}")
            return 1

    if not docx_only:
        for case in selected:
            generate_figures(case)

    if not png_only:
        for case in selected:
            generate_docx(case)

    if not png_only and not no_zip:
        extras = [
            "交代理人打包清单.md",
            "专利逻辑关系_一页讲清_曾欣宇.md",
        ]
        create_zip(selected, extras)

    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
