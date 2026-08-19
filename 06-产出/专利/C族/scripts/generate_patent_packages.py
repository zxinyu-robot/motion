#!/usr/bin/env python3
"""Batch-generate Word disclosure docs for C-family patents (PNG figures only)."""

from __future__ import annotations

import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent.parent
DISCLOSURE_DIR = BASE / "交底书"
OUT_DIR = BASE / "正式包"
FIG_DIR = BASE / "附图"


@dataclass
class PatentCase:
    key: str
    md_name: str
    figures: list[str] = field(default_factory=list)

    @property
    def md_path(self) -> Path:
        return DISCLOSURE_DIR / self.md_name

    @property
    def docx_path(self) -> Path:
        return OUT_DIR / Path(self.md_name).with_suffix(".docx").name


CASES: list[PatentCase] = [
    PatentCase(
        "Ca",
        "Ca_发明专利_端侧增量分层空间地图库_技术交底书.md",
        [
            "Ca-图1-端侧分层地图库总体结构.png",
            "Ca-图2-分层地图库管理流程.png",
            "Ca-图3-分层空间地图库逻辑结构.png",
            "Ca-图4-分层库多消费者投影.png",
        ],
    ),
    PatentCase(
        "Cb",
        "Cb_发明专利_内容寻址LayerDelta封装与版本组合_技术交底书.md",
        [
            "Cb-图1-LayerDelta与版本DAG总体结构.png",
            "Cb-图2-封装与compose流程.png",
            "Cb-图3-LayerDelta信封结构.png",
            "Cb-图4-版本DAG与compose.png",
        ],
    ),
    PatentCase(
        "Cc",
        "Cc_发明专利_AccessTicket消费契约分发_技术交底书.md",
        [
            "Cc-图1-AccessTicket分发总体结构.png",
            "Cc-图2-令牌校验与物化流程.png",
        ],
    ),
    PatentCase(
        "Ch",
        "Ch_发明专利_多机器人地图版本并发管理_技术交底书.md",
        [
            "Ch-图1-多机并发版本管理总体结构.png",
            "Ch-图2-OCC与冲突仲裁流程.png",
        ],
    ),
]


def figure_keys(index: int) -> list[str]:
    n = str(index)
    return [f"图{n}", f"图 {n}"]


def build_figure_map(png_names: list[str]) -> dict[str, Path]:
    mapping: dict[str, Path] = {}
    for i, name in enumerate(png_names, start=1):
        png = FIG_DIR / name
        for key in figure_keys(i):
            mapping[key] = png
    return mapping


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text.strip(), level=min(level, 3))
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = row[j] if j < len(row) else ""
            cell_text = re.sub(r"\*\*(.+?)\*\*", r"\1", cell_text)
            cell_text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", cell_text)
            cell_text = re.sub(r"`(.+?)`", r"\1", cell_text)
            table.rows[i].cells[j].text = cell_text


def figure_for_caption(line: str, figures: dict[str, Path]) -> Path | None:
    for key, path in figures.items():
        if key in line and path.exists():
            return path
    return None


def add_figure(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.5))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.font.size = Pt(10.5)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_paragraph(doc: Document, text: str) -> None:
    text = text.strip()
    if not text:
        return
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_bullet(doc: Document, text: str) -> None:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text.strip())
    text = re.sub(r"`(.+?)`", r"\1", text)
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_numbered(doc: Document, text: str) -> None:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text.strip())
    text = re.sub(r"`(.+?)`", r"\1", text)
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def build_docx(md_path: Path, figures: dict[str, Path], md_text: str) -> Document:
    doc = Document()
    set_doc_font(doc)

    max_fig = max(
        (int(re.search(r"\d+", k).group()) for k in figures if re.search(r"\d+", k)),
        default=4,
    )
    fig_pattern = re.compile(rf"^图\s*[1-{max_fig}]")

    def is_internal_metadata_heading(text: str) -> bool:
        if text.startswith("## 文档元数据"):
            return True
        return bool(re.match(r"^## \d+[、.]?\s*元数据", text))

    lines = md_text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if is_internal_metadata_heading(stripped):
            break

        if stripped.startswith(">"):
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("##"):
            title = stripped.lstrip("# ").strip()
            t = doc.add_paragraph()
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = t.add_run(title)
            run.bold = True
            run.font.size = Pt(16)
            run.font.name = "黑体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            i += 1
            continue

        if stripped.startswith("## "):
            add_heading(doc, stripped[3:], 1)
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 2)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table_rows(table_lines))
            continue

        if re.match(r"^\d+\.\s", stripped):
            add_numbered(doc, re.sub(r"^\d+\.\s*", "", stripped))
            i += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            i += 1
            continue

        if stripped.startswith("![") and "](" in stripped:
            m = re.match(r"!\[(.+?)\]\((.+?)\)", stripped)
            if m:
                caption, rel = m.group(1), m.group(2)
                img_path = (md_path.parent / rel).resolve()
                if img_path.exists():
                    add_figure(doc, img_path, caption)
            i += 1
            continue

        fig_path = figure_for_caption(stripped, figures)
        if fig_path and fig_pattern.match(stripped):
            add_figure(doc, fig_path, stripped)
            i += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and len(stripped) < 120:
            p = doc.add_paragraph(stripped.strip("*"))
            for run in p.runs:
                run.bold = True
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            i += 1
            continue

        if stripped:
            add_paragraph(doc, stripped)

        i += 1

    return doc


def generate_docx(case: PatentCase) -> Path:
    md_path = case.md_path
    if not md_path.exists():
        raise FileNotFoundError(f"Missing markdown: {md_path}")

    figures = build_figure_map(case.figures)
    for path in figures.values():
        if not path.exists():
            raise FileNotFoundError(f"Missing PNG figure: {path}")

    md_text = md_path.read_text(encoding="utf-8")
    doc = build_docx(md_path, figures, md_text)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(case.docx_path)
    print(f"DOCX: {case.docx_path.relative_to(BASE)}")
    return case.docx_path


def create_zip(cases: list[PatentCase], extra_files: list[tuple[str, str]]) -> Path:
    """extra_files: list of (disk_path_relative_to_BASE, arcname_in_zip)."""
    stamp = date.today().strftime("%Y%m%d")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    zip_path = OUT_DIR / f"C族-Cb核心-CcCh卫星-{stamp}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for case in cases:
            zf.write(case.docx_path, case.docx_path.name)
            for fig in case.figures:
                png = FIG_DIR / fig
                zf.write(png, f"附图/{png.name}")
        for rel_path, arcname in extra_files:
            path = BASE / rel_path
            if path.exists():
                zf.write(path, arcname)
    print(f"ZIP: {zip_path.relative_to(BASE)}")
    return zip_path


def main(argv: list[str]) -> int:
    only_keys = [a for a in argv if not a.startswith("-")]
    no_zip = "--no-zip" in argv

    selected = CASES
    if only_keys:
        selected = [c for c in CASES if c.key in only_keys]
        if not selected:
            print(f"Unknown keys: {only_keys}. Valid: {[c.key for c in CASES]}")
            return 1

    for case in selected:
        generate_docx(case)

    if not no_zip:
        extras = [
            ("交代理人打包清单.md", "交代理人打包清单.md"),
            ("专利逻辑关系_一页讲清_C族.md", "专利逻辑关系_一页讲清_C族.md"),
            ("叙事/C族技术交底故事.md", "C族技术交底故事.md"),
        ]
        create_zip(selected, extras)

    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
